from __future__ import annotations

import copy
import math
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CAPS_Major_Project_Phase1_Documentation_University_Logo.docx"
OUTPUT = ROOT / "CAPS_Major_Project_Phase1_Documentation_University_Logo_Formatted.docx"
ASSETS = ROOT / "research_pilot" / "docx_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

INK = "#263238"
NAVY = "#264653"
TEAL = "#2A9D8F"
GOLD = "#E9C46A"
CORAL = "#E76F51"
MINT = "#E6F2EF"
PALE_BLUE = "#EAF1F5"
PALE_GOLD = "#FBF3DB"
PALE_CORAL = "#F9E8E3"
WHITE = "#FFFFFF"
GRAY = "#68747A"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    filename = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / filename
    if path.exists():
        return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, face: ImageFont.ImageFont, max_width: int) -> str:
    result: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        line = ""
        for word in words:
            candidate = f"{line} {word}".strip()
            if line and draw.textlength(candidate, font=face) > max_width:
                result.append(line)
                line = word
            else:
                line = candidate
        result.append(line)
    return "\n".join(result)


def box(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], text: str,
        fill: str = PALE_BLUE, outline: str = NAVY, size: int = 31,
        bold: bool = True) -> None:
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=24, fill=fill, outline=outline, width=4)
    face = font(size, bold)
    wrapped = wrap_lines(draw, text, face, x2 - x1 - 36)
    bounds = draw.multiline_textbbox((0, 0), wrapped, font=face, spacing=10, align="center")
    text_w = bounds[2] - bounds[0]
    text_h = bounds[3] - bounds[1]
    text_x = (x1 + x2 - text_w) / 2 - bounds[0]
    text_y = (y1 + y2 - text_h) / 2 - bounds[1]
    draw.multiline_text((text_x, text_y), wrapped, font=face, fill=INK, spacing=10, align="center")


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
          color: str = NAVY, width: int = 8, head: int = 22) -> None:
    draw.line((start, end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    left = (end[0] - head * math.cos(angle - math.pi / 6),
            end[1] - head * math.sin(angle - math.pi / 6))
    right = (end[0] - head * math.cos(angle + math.pi / 6),
             end[1] - head * math.sin(angle + math.pi / 6))
    draw.polygon((end, left, right), fill=color)


def elbow_arrow(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]],
                color: str = TEAL, width: int = 8) -> None:
    draw.line(points, fill=color, width=width, joint="curve")
    arrow(draw, points[-2], points[-1], color=color, width=width)


def figure_canvas(title: str, subtitle: str, height: int = 1100) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1800, height), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((900, 60), title, font=font(48, True), fill=NAVY, anchor="mm")
    draw.text((900, 122), subtitle, font=font(27), fill=GRAY, anchor="mm")
    return image, draw


def make_figure_21(path: Path) -> None:
    image, draw = figure_canvas(
        "Communication strategies in a multi-agent workflow",
        "The policy changes which information reaches the receiving agent.",
        height=1040,
    )
    draw.text((80, 210), "Conventional full-message exchange", font=font(34, True), fill=INK)
    row1 = [(80, 285, 420, 465), (525, 285, 865, 465), (970, 285, 1310, 465), (1415, 285, 1755, 465)]
    labels1 = ["Sender agent", "Full message\nincluding repeated detail", "Receiver agent", "Higher context and\ncommunication cost"]
    fills1 = [PALE_BLUE, PALE_CORAL, PALE_BLUE, PALE_GOLD]
    for rect, label, fill in zip(row1, labels1, fills1):
        box(draw, rect, label, fill=fill, outline=NAVY if fill == PALE_BLUE else CORAL if fill == PALE_CORAL else "#B18A20")
    for a, b in zip(row1, row1[1:]):
        arrow(draw, (a[2] + 15, 375), (b[0] - 15, 375))

    draw.text((80, 570), "CAPS selective communication", font=font(34, True), fill=INK)
    row2 = [(70, 650, 360, 845), (420, 650, 710, 845), (770, 650, 1060, 845),
            (1120, 650, 1410, 845), (1470, 650, 1760, 845)]
    labels2 = ["Sender agent", "Extract candidate\ninformation units", "Policy weighs context,\nvalue, and token cost",
               "Transmit selected\nunits within budget", "Receiver agent"]
    fills2 = [PALE_BLUE, MINT, PALE_GOLD, MINT, PALE_BLUE]
    borders2 = [NAVY, TEAL, "#B18A20", TEAL, NAVY]
    for rect, label, fill, border in zip(row2, labels2, fills2, borders2):
        box(draw, rect, label, fill=fill, outline=border, size=29)
    for a, b in zip(row2, row2[1:]):
        arrow(draw, (a[2] + 12, 747), (b[0] - 12, 747))
    draw.text((900, 950), "Selective transmission aims to preserve task-relevant information while limiting unnecessary context.",
              font=font(25), fill=GRAY, anchor="mm")
    image.save(path, dpi=(300, 300))


def make_figure_51(path: Path) -> None:
    image, draw = figure_canvas(
        "Proposed CAPS system architecture",
        "A policy layer selects agent communication and learns from task outcomes.",
        height=1120,
    )
    nodes = [
        (45, 270, 275, 485, "Task and\nworkflow", PALE_BLUE, NAVY),
        (295, 270, 525, 485, "Agent\nadapter", PALE_BLUE, NAVY),
        (545, 270, 775, 485, "CIU\nextractor", MINT, TEAL),
        (795, 270, 1025, 485, "Context\nbuilder", MINT, TEAL),
        (1045, 270, 1275, 485, "Policy\nselector", PALE_GOLD, "#B18A20"),
        (1295, 270, 1525, 485, "Communication\ncontroller", PALE_GOLD, "#B18A20"),
        (1545, 270, 1775, 485, "Receiving\nagent", PALE_BLUE, NAVY),
    ]
    for x1, y1, x2, y2, label, fill, border in nodes:
        box(draw, (x1, y1, x2, y2), label, fill=fill, outline=border, size=28)
    for a, b in zip(nodes, nodes[1:]):
        arrow(draw, (a[2] + 7, 377), (b[0] - 7, 377), width=7, head=18)

    box(draw, (330, 690, 700, 895), "Trace logger\n(decisions, usage, outcomes)", fill=PALE_BLUE, outline=NAVY, size=28)
    box(draw, (790, 690, 1160, 895), "Task evaluator\n(quality and constraints)", fill=PALE_CORAL, outline=CORAL, size=28)
    box(draw, (1250, 690, 1620, 895), "Reward signal\n(quality relative to cost)", fill=PALE_GOLD, outline="#B18A20", size=28)
    elbow_arrow(draw, [(1650, 485), (1650, 620), (975, 620), (975, 680)], color=CORAL, width=7)
    arrow(draw, (1170, 792), (1235, 792), color=CORAL, width=7, head=20)
    elbow_arrow(draw, [(1435, 680), (1435, 580), (1160, 580), (1160, 495)], color=TEAL, width=7)
    draw.text((900, 1010), "The logger records each decision; observed task outcomes update the policy during training.",
              font=font(25), fill=GRAY, anchor="mm")
    image.save(path, dpi=(300, 300))


def make_figure_52(path: Path) -> None:
    image, draw = figure_canvas(
        "End-to-end policy learning workflow",
        "Each episode produces a communication decision and an observable task outcome.",
        height=1080,
    )
    nodes = [
        (55, 285, 315, 500, "Task and\ncollaboration state", PALE_BLUE, NAVY),
        (350, 285, 610, 500, "Agents produce\ncandidate messages", PALE_BLUE, NAVY),
        (645, 285, 905, 500, "Extract CIUs and\nestimate costs", MINT, TEAL),
        (940, 285, 1200, 500, "Policy scores\navailable actions", PALE_GOLD, "#B18A20"),
        (1235, 285, 1495, 500, "Deliver selected\ninformation", MINT, TEAL),
        (1530, 285, 1790, 500, "Evaluate task\noutcome", PALE_CORAL, CORAL),
    ]
    for x1, y1, x2, y2, label, fill, border in nodes:
        box(draw, (x1, y1, x2, y2), label, fill=fill, outline=border, size=27)
    for a, b in zip(nodes, nodes[1:]):
        arrow(draw, (a[2] + 7, 392), (b[0] - 7, 392), width=7, head=18)
    box(draw, (650, 690, 1150, 890), "Update contextual policy\nfrom observed reward", fill=PALE_GOLD, outline="#B18A20", size=31)
    elbow_arrow(draw, [(1660, 500), (1660, 790), (1170, 790)], color=CORAL, width=7)
    elbow_arrow(draw, [(650, 790), (480, 790), (480, 600), (1070, 600), (1070, 515)], color=TEAL, width=7)
    draw.text((900, 970), "Policy updates use training feedback; validation and held-out test policies remain frozen.",
              font=font(25), fill=GRAY, anchor="mm")
    image.save(path, dpi=(300, 300))


def make_figure_53(path: Path) -> None:
    image, draw = figure_canvas(
        "Evaluation and feedback loop",
        "Development informs policy selection; held-out results are reserved for final reporting.",
        height=1210,
    )
    draw.text((75, 205), "Development", font=font(34, True), fill=INK)
    top = [
        (70, 270, 440, 485, "Training tasks\nfit the policy", PALE_BLUE, NAVY),
        (500, 270, 870, 485, "Validation tasks\nselect settings", PALE_GOLD, "#B18A20"),
        (930, 270, 1300, 485, "Freeze policy,\nfeatures, and budget", MINT, TEAL),
    ]
    for x1, y1, x2, y2, label, fill, border in top:
        box(draw, (x1, y1, x2, y2), label, fill=fill, outline=border, size=29)
    for a, b in zip(top, top[1:]):
        arrow(draw, (a[2] + 12, 378), (b[0] - 12, 378))
    arrow(draw, (685, 500), (685, 610), color="#B18A20", width=6, head=18)
    arrow(draw, (685, 610), (255, 610), color="#B18A20", width=6, head=18)
    arrow(draw, (255, 610), (255, 500), color="#B18A20", width=6, head=18)
    draw.text((470, 570), "Validation may guide configuration", font=font(23), fill=GRAY, anchor="mm")

    draw.text((75, 690), "Held-out evaluation", font=font(34, True), fill=INK)
    bottom = [
        (55, 760, 375, 970, "Held-out task\ninstances", PALE_BLUE, NAVY),
        (410, 760, 730, 970, "Matched runs for\neach baseline", PALE_BLUE, NAVY),
        (765, 760, 1085, 970, "Task success and\nresource measures", MINT, TEAL),
        (1120, 760, 1440, 970, "Paired statistics\nand uncertainty", PALE_GOLD, "#B18A20"),
        (1475, 760, 1795, 970, "Report and\nerror analysis", PALE_CORAL, CORAL),
    ]
    for x1, y1, x2, y2, label, fill, border in bottom:
        box(draw, (x1, y1, x2, y2), label, fill=fill, outline=border, size=27)
    for a, b in zip(bottom, bottom[1:]):
        arrow(draw, (a[2] + 8, 865), (b[0] - 8, 865), width=7, head=18)
    elbow_arrow(draw, [(1115, 485), (1115, 650), (570, 650), (570, 750)], color=TEAL, width=7)
    draw.text((900, 1090), "Test outcomes inform the report only; they do not feed back into policy or hyperparameter selection.",
              font=font(24), fill=GRAY, anchor="mm")
    image.save(path, dpi=(300, 300))


def draw_field(run, instruction: str, cached: str = "1") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = cached
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, code, separate, value, end))


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    p = paragraph._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insert_at = 1 if p.pPr is not None else 0
    p.insert(insert_at, start)
    p.append(end)


def set_page_numbering(sect_pr, fmt: str, start: int = 1) -> None:
    old = sect_pr.find(qn("w:pgNumType"))
    if old is not None:
        sect_pr.remove(old)
    item = OxmlElement("w:pgNumType")
    item.set(qn("w:fmt"), fmt)
    item.set(qn("w:start"), str(start))
    sect_pr.append(item)


def add_footer_field(section) -> None:
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    draw_field(run, "PAGE")


def add_picture_before(doc, anchor, path: Path, alt: str, caption_text: str | None = None):
    image_para = doc.add_paragraph()
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.paragraph_format.space_before = Pt(4)
    image_para.paragraph_format.space_after = Pt(2)
    image_para.paragraph_format.keep_with_next = True
    picture = image_para.add_run().add_picture(str(path), width=Inches(5.25))
    picture._inline.docPr.set("descr", alt)
    if caption_text:
        picture._inline.docPr.set("title", caption_text)
    anchor._p.addprevious(image_para._p)
    return image_para


def add_caption_before(doc, anchor, text: str):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.keep_together = True
    try:
        paragraph.style = "Caption"
    except KeyError:
        pass
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = False
    anchor._p.addprevious(paragraph._p)
    return paragraph


def reset_caption(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.keep_together = True
    try:
        paragraph.style = "Caption"
    except KeyError:
        pass
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = False


def add_candidate_actions_table(doc):
    anchor = next(
        p for p in doc.paragraphs
        if p.text.strip().startswith("For a linear contextual bandit, the estimated reward")
    )
    caption = doc.add_paragraph()
    caption.add_run("Table 5.2 Candidate policy actions")
    caption.style = "Caption"
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    for run in caption.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = False

    rows = [
        ("Action", "Communication behavior"),
        ("A1 Full message", "Transmit the complete candidate message."),
        ("A2 Summary", "Transmit a concise summary of the candidate message."),
        ("A3 Selected CIUs", "Transmit only the selected contextual information units."),
        ("A4 Defer", "Hold communication until a defined checkpoint."),
        ("A5 Suppress", "Withhold a candidate message when it is redundant and nonessential."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(3.95)
    for row_index, values in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = value
            cell.width = Inches(1.55 if col_index == 0 else 3.95)
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
            if row_index == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E2F3")
                tc_pr.append(shading)

    anchor._p.addprevious(table._tbl)
    anchor._p.addprevious(caption._p)


def replace_with_page_ref(paragraph, label: str, bookmark_name: str,
                          content_width, indent_inches: float = 0) -> None:
    paragraph.clear()
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(indent_inches)
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.clear_all()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        content_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    paragraph.add_run(label)
    paragraph.add_run().add_tab()
    field_run = paragraph.add_run()
    field_run.font.name = "Times New Roman"
    field_run.font.size = Pt(11)
    draw_field(field_run, f"PAGEREF {bookmark_name} \\h")


def find_paragraph(doc, exact_text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    key = " ".join(exact_text.split()).casefold()
    for paragraph in doc.paragraphs:
        if " ".join(paragraph.text.split()).casefold() == key:
            return paragraph
    raise ValueError(f"Required paragraph not found: {exact_text}")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    figure_paths = {
        "2.1": ASSETS / "figure_2_1_communication_strategies.png",
        "5.1": ASSETS / "figure_5_1_caps_architecture.png",
        "5.2": ASSETS / "figure_5_2_policy_workflow.png",
        "5.3": ASSETS / "figure_5_3_evaluation_loop.png",
    }
    make_figure_21(figure_paths["2.1"])
    make_figure_51(figure_paths["5.1"])
    make_figure_52(figure_paths["5.2"])
    make_figure_53(figure_paths["5.3"])

    doc = Document(SOURCE)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    # Restore Chapter 1 prose to the document's Normal style and promote its numbered headings.
    chapter2_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "CHAPTER 2")
    intro_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "1.1 Introduction")
    for paragraph in doc.paragraphs[intro_index:chapter2_index]:
        text = paragraph.text.strip()
        if not text:
            continue
        heading_match = re.match(r"^1\.\d+(?:\.\d+)?\s+", text)
        if heading_match:
            paragraph.style = "Heading 3" if text.startswith("1.5.1 ") else "Heading 2"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = None
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_before = None
            paragraph.paragraph_format.space_after = None
            paragraph.paragraph_format.line_spacing = None
            for run in paragraph.runs:
                run.bold = None
                run.font.size = None
                run.font.name = None
            continue

        original_runs = list(paragraph.runs)
        paragraph.style = "Normal"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = None
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = None
        paragraph.paragraph_format.space_after = None
        paragraph.paragraph_format.line_spacing = None
        for index, run in enumerate(original_runs):
            run.bold = True if index == 0 and len(original_runs) > 1 and run.text.rstrip().endswith(":") else None
            run.font.size = None
            run.font.name = None

    # Add a new-page section boundary immediately before Chapter 1.
    chapter1_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "CHAPTER 1")
    preceding = doc.paragraphs[chapter1_index - 1]
    body_sect_pr = doc._element.body.sectPr
    first_sect_pr = copy.deepcopy(body_sect_pr)
    section_type = first_sect_pr.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        first_sect_pr.insert(0, section_type)
    section_type.set(qn("w:val"), "nextPage")
    p_pr = preceding._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:sectPr"))
    if existing is not None:
        p_pr.remove(existing)
    p_pr.append(first_sect_pr)

    sections = doc.sections
    if len(sections) != 2:
        raise RuntimeError(f"Expected two sections after split, got {len(sections)}")
    set_page_numbering(sections[0]._sectPr, "upperRoman", 1)
    set_page_numbering(sections[1]._sectPr, "decimal", 1)
    sections[0].different_first_page_header_footer = True
    sections[1].different_first_page_header_footer = False
    sections[0].first_page_footer.paragraphs[0].clear()
    add_footer_field(sections[0])
    add_footer_field(sections[1])

    # Replace incomplete figure notes with four legible, captioned research diagrams.
    figure_21_caption = find_paragraph(doc, "Figure 2.1. Existing and proposed communication approaches (schematic to be inserted in the final formatted copy).")
    reset_caption(figure_21_caption, "Figure 2.1. Existing and proposed communication approaches")
    add_picture_before(doc, figure_21_caption, figure_paths["2.1"],
                       "Comparison of full-message communication with CAPS selective transmission of contextual information units.",
                       "Figure 2.1. Existing and proposed communication approaches")

    figure_51_caption = find_paragraph(doc, "Figure 5.1. Proposed CAPS system architecture (insert architecture diagram during final formatting).")
    reset_caption(figure_51_caption, "Figure 5.1. Proposed CAPS system architecture")
    add_picture_before(doc, figure_51_caption, figure_paths["5.1"],
                       "CAPS architecture from task orchestration through CIU extraction, context building, policy selection, communication control, evaluation, logging, and reward feedback.",
                       "Figure 5.1. Proposed CAPS system architecture")

    heading_52 = find_paragraph(doc, "5.2 Contextual Information Units")
    add_picture_before(doc, heading_52, figure_paths["5.2"],
                       "Policy-learning workflow from task context through message generation, CIU extraction, policy selection, delivery, outcome evaluation, and training update.",
                       "Figure 5.2. End-to-end policy learning workflow")
    add_caption_before(doc, heading_52, "Figure 5.2. End-to-end policy learning workflow")

    heading_57 = find_paragraph(doc, "5.7 Evaluation Metrics")
    add_picture_before(doc, heading_57, figure_paths["5.3"],
                       "Evaluation workflow showing training and validation for policy selection, followed by frozen matched-budget testing and statistical reporting.",
                       "Figure 5.3. Proposed evaluation and feedback loop")
    add_caption_before(doc, heading_57, "Figure 5.3. Proposed evaluation and feedback loop")

    add_candidate_actions_table(doc)

    metadata_heading = find_paragraph(doc, "A.1 Run-level metadata")
    metadata_heading.paragraph_format.page_break_before = True

    # Add bookmarks to every destination and convert static contents pages to live references.
    bookmark_id = 1
    bookmark_names: dict[str, str] = {}

    def mark(paragraph, name: str) -> str:
        nonlocal bookmark_id
        safe_name = re.sub(r"[^A-Za-z0-9_]", "_", name)
        add_bookmark(paragraph, safe_name, bookmark_id)
        bookmark_id += 1
        return safe_name

    content_width = sections[0].page_width - sections[0].left_margin - sections[0].right_margin
    toc_targets = {
        "chapter 1 introduction": "CHAPTER 1",
        "chapter 2 problem identification": "CHAPTER 2",
        "chapter 3 literature survey": "CHAPTER 3",
        "chapter 4 feasibility study and requirements": "CHAPTER 4",
        "chapter 5 proposed methodology": "CHAPTER 5",
        "chapter 6 expected outcomes and phase-1 conclusion": "CHAPTER 6",
        "appendix a proposed experiment matrix": "APPENDIX A",
        "appendix b traceability and submission checklist": "APPENDIX B",
    }
    toc_heading = find_paragraph(doc, "TABLE OF CONTENTS")
    lof_heading = find_paragraph(doc, "LIST OF FIGURES")
    lot_heading = find_paragraph(doc, "LIST OF TABLES")
    body_paragraphs = doc.paragraphs
    toc_start = next(i for i, p in enumerate(body_paragraphs) if p.text.strip() == "TABLE OF CONTENTS") + 1
    lof_start = next(i for i, p in enumerate(body_paragraphs) if p.text.strip() == "LIST OF FIGURES")
    for paragraph in body_paragraphs[toc_start:lof_start]:
        original = paragraph.text.strip()
        if "\t" not in original:
            continue
        label = original.rsplit("\t", 1)[0].strip()
        target = toc_targets.get(label.casefold(), label.upper())
        target_paragraph = find_paragraph(doc, target)
        bookmark = bookmark_names.get(target)
        if bookmark is None:
            bookmark = mark(target_paragraph, "toc_" + re.sub(r"\W+", "_", label.lower()))
            bookmark_names[target] = bookmark
        indent = 0.28 if re.match(r"^[1-6]\.\d+", label) else 0
        replace_with_page_ref(paragraph, label, bookmark, content_width, indent)

    figure_captions = {
        "Figure 2.1": find_paragraph(doc, "Figure 2.1. Existing and proposed communication approaches"),
        "Figure 5.1": find_paragraph(doc, "Figure 5.1. Proposed CAPS system architecture"),
        "Figure 5.2": find_paragraph(doc, "Figure 5.2. End-to-end policy learning workflow"),
        "Figure 5.3": find_paragraph(doc, "Figure 5.3. Proposed evaluation and feedback loop"),
    }
    for key, paragraph in figure_captions.items():
        name = "fig_" + key.replace(" ", "_").replace(".", "_")
        bookmark_names[key] = mark(paragraph, name)

    figure_list_start = next(i for i, p in enumerate(body_paragraphs) if p.text.strip() == "LIST OF FIGURES") + 1
    table_list_start = next(i for i, p in enumerate(body_paragraphs) if p.text.strip() == "LIST OF TABLES")
    figure_label_map = {
        "Figure 2.1": "Figure 2.1",
        "Figure 5.1": "Figure 5.1",
        "Figure 5.2": "Figure 5.2",
        "Figure 5.3": "Figure 5.3",
    }
    for paragraph in body_paragraphs[figure_list_start:table_list_start]:
        original = paragraph.text.strip()
        match = re.match(r"(Figure\s+\d+\.\d+)", original)
        if not match:
            continue
        key = figure_label_map[match.group(1)]
        label = original.split("[PAGE PLACEHOLDER]")[0].rstrip(" .")
        replace_with_page_ref(paragraph, label, bookmark_names[key], content_width)

    abbreviations_index = next(i for i, p in enumerate(body_paragraphs)
                               if p.text.strip() == "ABBREVIATIONS AND NOTATION")
    table_entries = [p for p in body_paragraphs[table_list_start + 1:abbreviations_index]
                     if re.match(r"Table\s+(?:\d+\.\d+|A\.\d+)", p.text.strip())]
    if len(table_entries) != 12 or len(doc.tables) < 15:
        raise RuntimeError(f"Unexpected table listing: {len(table_entries)} entries, {len(doc.tables)} document tables")
    table_index_by_number = {
        "2.1": 2,
        "3.1": 3,
        "4.1": 4,
        "4.2": 6,
        "4.3": 7,
        "4.4": 8,
        "5.1": 9,
        "5.2": 10,
        "5.3": 12,
        "5.4": 11,
        "6.1": 13,
        "A.1": 14,
    }
    for paragraph in table_entries:
        label = paragraph.text.split("[PAGE PLACEHOLDER]")[0].rstrip(" .")
        number = re.match(r"Table\s+((?:\d+\.\d+)|A\.\d+)", label).group(1)
        table = doc.tables[table_index_by_number[number]]
        anchor_paragraph = table.cell(0, 0).paragraphs[0]
        bookmark = mark(anchor_paragraph, f"table_ref_{number.replace('.', '_')}")
        replace_with_page_ref(paragraph, label, bookmark, content_width)

    for table in doc.tables:
        if not table.rows:
            continue
        for row_index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0 and row_properties.find(qn("w:tblHeader")) is None:
                repeated = OxmlElement("w:tblHeader")
                repeated.set(qn("w:val"), "true")
                row_properties.append(repeated)

    # Ask Word to refresh PAGE/PAGEREF fields when the file is opened.
    update = doc.settings.element.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        doc.settings.element.append(update)
    update.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")
    print(f"Figures: {len(figure_paths)}; sections: {len(doc.sections)}; tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
